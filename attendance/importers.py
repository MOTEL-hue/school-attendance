"""ייבוא תלמידות מאקסל/CSV וייצוא תבנית.

עמודות: ת.ז. | שם פרטי | שם משפחה | כיתה | טלפון אמא | טלפון אבא | טלפון תלמידה
"""
import csv
import io

from openpyxl import Workbook, load_workbook

from .models import Contact, SchoolClass, Student, clean_phone

HEADERS = ["ת.ז.", "שם פרטי", "שם משפחה", "כיתה", "טלפון אמא", "טלפון אבא", "טלפון תלמידה"]
CONTACT_COLS = {4: "אמא", 5: "אבא", 6: "תלמידה"}


def template_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append(["012345678", "שרה", "כהן", "ח1", "0501234567", "", ""])
    ws.sheet_view.rightToLeft = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_rows(f):
    name = f.name.lower()
    if name.endswith(".xlsx"):
        ws = load_workbook(f, read_only=True, data_only=True).active
        return [[("" if c is None else str(c).strip()) for c in row] for row in ws.iter_rows(values_only=True)]
    raw = f.read()
    for enc in ("utf-8-sig", "cp1255"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("לא ניתן לקרוא את הקובץ (קידוד)")
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]


def import_students(f):
    """מחזיר (נוספו, עודכנו, רשימת שגיאות)."""
    rows = read_rows(f)
    if rows and not (rows[0] and rows[0][0].replace(" ", "").isdigit()):
        rows = rows[1:]  # שורת כותרת
    added = updated = 0
    errors = []
    for n, row in enumerate(rows, start=2):
        row = (row + [""] * 7)[:7]
        if not any(row):
            continue
        tz = "".join(c for c in row[0] if c.isdigit())
        if not (5 <= len(tz) <= 9) or not row[1] or not row[3]:
            errors.append(f"שורה {n}: חסר או לא תקין (ת.ז. / שם / כיתה)")
            continue
        cls, _ = SchoolClass.objects.get_or_create(name=row[3])
        st, created = Student.objects.update_or_create(
            tz=tz.zfill(9), defaults={"first_name": row[1], "last_name": row[2], "school_class": cls, "active": True})
        added += created
        updated += not created
        for col, label in CONTACT_COLS.items():
            phone = clean_phone(row[col])
            if phone:
                Contact.objects.get_or_create(student=st, phone=phone, defaults={"label": label})
    return added, updated, errors
